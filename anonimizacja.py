#!/usr/bin/env python3
"""Anonimizacja RODO. Wyszukuje dane osobowe w DOCX, PDF, skanach i zdjeciach
PNG/JPG (PESEL, NIP, REGON, nr dowodu, IBAN, e-mail, telefon + wlasna lista
slow) i trwale je usuwa: w PDF tekst jest wycinany z pliku (nie tylko
zakrywany), w DOCX zamieniany na gwiazdki, na skanach i zdjeciach (OCR,
Tesseract) zamalowywany na czarno w pikselach. Czysci tez metadane (autor,
GPS ze zdjec itp.).

Frazy do pominiecia (np. NIP urzedu) zostaja w dokumencie. Skanowanie
(skan=True) tylko zaznacza dane w kopii _skan, nic nie usuwa. Po kazdym
przebiegu log wypisuje liste znalezionych danych.

Uruchomienie: python anonimizacja.py            (GUI)
              python anonimizacja.py --selftest (test logiki)
"""
import csv
import io
import os
import re
import shutil
import subprocess
import sys
import threading
import traceback
from collections import Counter

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))
# wlasny model polski (tessdata_best) - Tesseract z instalatora ma tylko angielski
TESSDATA_DIR = os.path.join(APP_DIR, "tessdata")
OCR_DPI = 300
IMAGE_EXTS = (".png", ".jpg", ".jpeg")
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)

# ------------------------------------------------------------ walidatory ----


def _digits(s):
    return [int(c) for c in s if c.isdigit()]


def _weighted(d, weights):
    return sum(a * b for a, b in zip(d, weights))


def pesel_ok(s):
    d = _digits(s)
    return len(d) == 11 and (10 - _weighted(d[:10], [1, 3, 7, 9] * 3) % 10) % 10 == d[10]


def nip_ok(s):
    d = _digits(s)
    return len(d) == 10 and _weighted(d, [6, 5, 7, 2, 3, 4, 5, 6, 7]) % 11 == d[9]


def regon_ok(s):
    d = _digits(s)
    if len(d) == 9:
        w = [8, 9, 2, 3, 4, 5, 6, 7]
    elif len(d) == 14:
        w = [2, 4, 8, 5, 0, 9, 7, 3, 6, 1, 2, 4, 8]
    else:
        return False
    return _weighted(d, w) % 11 % 10 == d[-1]


def dowod_ok(s):
    s = s.replace(" ", "").upper()
    vals = [ord(c) - 55 if c.isalpha() else int(c) for c in s]
    return len(vals) == 9 and _weighted(vals, [7, 3, 1, 9, 7, 3, 1, 7, 3]) % 10 == 0


def iban_ok(s):
    s = re.sub(r"\s", "", s).upper()
    if not s.startswith("PL"):
        s = "PL" + s
    if len(s) != 28:
        return False
    num = "".join(str(int(c, 36)) for c in s[4:] + s[:4])
    return int(num) % 97 == 1


# Kolejnosc = priorytet: dluzsze/pewniejsze wzorce zajmuja tekst pierwsze.
# (nazwa, regex, walidator lub None)
DETECTORS = [
    ("IBAN", r"(?<![\w])(?:PL ?)?\d{2}(?: ?\d{4}){6}(?!\d)", iban_ok),
    ("REGON", r"(?<![\d-])\d{14}(?![\d-])", regon_ok),
    ("PESEL", r"(?<![\d-])\d{11}(?![\d-])", pesel_ok),
    ("NIP", r"(?<![\w-])(?:PL ?)?(?:\d{3}-?\d{3}-?\d{2}-?\d{2}|\d{3}-\d{2}-\d{2}-\d{3})(?![\d-])",
     nip_ok),
    ("REGON", r"(?<![\d-])\d{9}(?![\d-])", regon_ok),
    ("Dowód", r"(?<![\w])[A-Z]{3} ?\d{6}(?!\d)", dowod_ok),
    ("E-mail", r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", None),
    ("Telefon", r"(?<![\d+])(?:(?:\+48|0048) ?)?(?:\d{3}[ -]\d{3}[ -]\d{3}|\(?\d{2}\)?[ -]\d{3}[ -]\d{2}[ -]\d{2})(?!\d)"
     r"|(?:\+48|0048) ?\d{9}(?!\d)", None),
]
CATEGORIES = list(dict.fromkeys(name for name, _, _ in DETECTORS))
# OCR potrafi pomylic jedna cyfre i wtedy suma kontrolna sie nie zgadza -
# na skanach lepiej zakryc za duzo niz przepuscic PESEL
OCR_EXTRA = [("PESEL (niepewny OCR)", r"(?<![\d-])\d{11}(?![\d-])", None)]
# Word i PDF-y uzywaja twardych spacji i roznych dywizow - znak na znak
# (dlugosc tekstu bez zmian, wiec pozycje trafien pasuja do oryginalu)
NORMALIZE = str.maketrans("   ­‐‑‒–−",
                          "   ------")


def _bez_separatorow(s):
    return re.sub(r"[\W_]", "", s.translate(NORMALIZE)).lower()


def _fraza(w):
    """Fraza jako regex: cale slowa, dowolne odstepy (PDF lamie linie), bez wielkosci liter."""
    return r"(?<!\w)%s(?!\w)" % r"\s+".join(map(re.escape, w.split()))


def find_spans(text, enabled=None, words=(), ocr=False, ignore=()):
    """Zwraca [(start, koniec, kategoria)] bez nakladania, posortowane.
    ocr=True: tekst z OCR - dodatkowo ciagi 11 cyfr z bledna suma kontrolna.
    ignore: frazy do pominiecia (np. NIP urzedu) - fragmenty tekstu z ta fraza
    i trafienia rowne frazie po usunieciu separatorow nie sa zwracane."""
    text = text.translate(NORMALIZE)
    taken = []
    ignore = {w.strip() for w in ignore if w.strip()}
    pomin = {_bez_separatorow(w) for w in ignore}

    def free(a, b):
        return all(b <= x or a >= y for x, y, _ in taken)

    # frazy pominiete zajmuja tekst pierwsze - zaden wzorzec nie wejdzie w ich srodek
    for w in ignore:
        for m in re.finditer(_fraza(w), text, re.I):
            if free(m.start(), m.end()):
                taken.append((m.start(), m.end(), None))
    for name, rx, check in DETECTORS + (OCR_EXTRA if ocr else []):
        if enabled is not None and name.split(" (")[0] not in enabled:
            continue
        for m in re.finditer(rx, text):
            if (check is None or check(m.group())) and free(m.start(), m.end()):
                # ten sam numer zapisany inaczej (123-456-32-18 / 1234563218)
                pominiety = _bez_separatorow(m.group()) in pomin
                taken.append((m.start(), m.end(), None if pominiety else name))
    for w in sorted({w.strip() for w in words if w.strip()}, key=len, reverse=True):
        for m in re.finditer(_fraza(w), text, re.I):
            if free(m.start(), m.end()):
                taken.append((m.start(), m.end(), "Lista słów"))
    return sorted(t for t in taken if t[2])


def znalezione(text, spans):
    """[(kategoria, tekst)] do raportu - odstepy i przelamania linii scalone."""
    return [(c, " ".join(text[a:b].split())) for a, b, c in spans]


def kategorie(found):
    """Counter((kategoria, tekst)) -> Counter(kategoria)."""
    out = Counter()
    for (c, _), n in found.items():
        out[c] += n
    return out


def mask(text, spans):
    """Zamienia znaki w zakresach na '*' (dlugosc tekstu bez zmian)."""
    chars = list(text)
    for a, b, _ in spans:
        for i in range(a, b):
            if not chars[i].isspace():
                chars[i] = "*"
    return "".join(chars)


# ------------------------------------------------------------------- OCR ----


def find_tesseract():
    """Silnik OCR (Tesseract) albo None. Instalator bez praw administratora
    wrzuca go do folderu uzytkownika, stad kilka lokalizacji."""
    local = os.environ.get("LOCALAPPDATA", "")
    for cand in (shutil.which("tesseract"),
                 r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                 r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                 os.path.join(local, "Programs", "Tesseract-OCR", "tesseract.exe"),
                 os.path.join(local, "Tesseract-OCR", "tesseract.exe")):
        if cand and os.path.exists(cand):
            return cand
    return None


def ocr_hits(png_bytes, enabled, words, ignore=()):
    """OCR obrazu -> [(kategoria, tekst, [prostokaty w pikselach])] dla znalezionych danych."""
    tess = find_tesseract()
    if not tess:
        raise ValueError("skan/zdjecie wymaga silnika OCR (Tesseract) - uruchom install.bat")
    # TESSDATA_PREFIX, nie --tessdata-dir (w Tesseract 5 psuje config 'tsv')
    out = subprocess.run([tess, "-", "-", "-l", "pol", "--dpi", str(OCR_DPI), "tsv"],
                         input=png_bytes, capture_output=True, check=True,
                         env={**os.environ, "TESSDATA_PREFIX": TESSDATA_DIR},
                         creationflags=NO_WINDOW).stdout.decode("utf8")
    paras = {}
    for r in csv.DictReader(io.StringIO(out), delimiter="\t", quoting=csv.QUOTE_NONE):
        if r["level"] == "5" and r["text"].strip():
            x, y, w, h = (int(r[k]) for k in ("left", "top", "width", "height"))
            pad = max(2, h // 6)  # brzegi liter nie moga wystawac spod prostokata
            paras.setdefault((r["page_num"], r["block_num"], r["par_num"]), []).append(
                (r["text"], (x - pad, y - pad, x + w + pad, y + h + pad)))
    hits = []
    # caly akapit naraz: IBAN czy telefon bywa przelamany miedzy liniami
    for ws in paras.values():
        text, starts = "", []
        for t, _ in ws:
            starts.append(len(text))
            text += t + " "
        for a, b, cat in find_spans(text, enabled, words, ocr=True, ignore=ignore):
            hits.append((cat, " ".join(text[a:b].split()),
                         [box for (t, box), s in zip(ws, starts) if s < b and s + len(t) > a]))
    return hits


def ramka(pix, b, kolor=(255, 0, 0), grubosc=3):
    """Skanowanie: czerwona ramka wokol prostokata (tresc pod spodem zostaje)."""
    import pymupdf
    x0, y0, x1, y1 = b
    g = grubosc
    for r in ((x0, y0, x1, y0 + g), (x0, y1 - g, x1, y1), (x0, y0, x0 + g, y1), (x1 - g, y0, x1, y1)):
        pix.set_rect(pymupdf.IRect(r) & pix.irect, kolor)


def anon_image(src, dst, enabled, words, ignore=(), skan=False):
    import pymupdf

    with pymupdf.open(src) as doc:  # obraz otwiera sie juz obrocony wg EXIF
        page = doc[0]
        orig = pymupdf.Pixmap(src)
        zoom = max(orig.width, orig.height) / max(page.rect.width, page.rect.height)
        pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
    hits = ocr_hits(pix.tobytes("png"), enabled, words, ignore)
    for _, _, boxes in hits:
        for b in boxes:
            if skan:
                ramka(pix, b)
            else:
                pix.set_rect(pymupdf.IRect(b) & pix.irect, (0, 0, 0))
    pix.set_dpi(orig.xres, orig.yres)
    # zapis od nowa z samych pikseli: bez EXIF (GPS, aparat, data), XMP i miniatury
    pix.save(dst, jpg_quality=95)
    return Counter((c, t) for c, t, _ in hits), 0


# ------------------------------------------------------------------ DOCX ----


def zaznacz_docx(t, segmenty):
    """Skanowanie: dzieli run z tekstem t na runy wg [(tekst, zaznaczony)],
    zaznaczone dostaja zolty marker. Formatowanie runu zostaje."""
    import copy
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.text.run import Run

    r = t.getparent()
    if r.tag != qn("w:r"):
        return
    rpr = r.find(qn("w:rPr"))

    def nowy_run():
        nr = OxmlElement("w:r")
        if rpr is not None:
            nr.append(copy.deepcopy(rpr))
        return nr

    nowe = []
    for tekst, zaznaczony in segmenty:
        nr = nowy_run()
        nt = OxmlElement("w:t")
        nt.text = tekst
        nt.set(qn("xml:space"), "preserve")
        nr.append(nt)
        if zaznaczony:
            Run(nr, None).font.highlight_color = WD_COLOR_INDEX.YELLOW
        nowe.append(nr)
    po = list(t.itersiblings())  # np. tabulator za tekstem - zostaje za nim
    if po:
        nr = nowy_run()
        nr.extend(po)
        nowe.append(nr)
    r.remove(t)
    for n in reversed(nowe):
        r.addnext(n)
    if all(c.tag == qn("w:rPr") for c in r):
        r.getparent().remove(r)


def anon_docx(src, dst, enabled, words, ignore=(), skan=False):
    import docx
    from docx.oxml.ns import qn

    doc = docx.Document(src)
    parts = [doc.part] + [r.target_part for r in doc.part.rels.values()
                          if r.reltype.endswith(("/header", "/footer",
                                                 "/footnotes", "/endnotes"))]
    found = Counter()
    # delText = usuniete w trybie sledzenia zmian, instrText = pole HYPERLINK "mailto:..."
    tags = (qn("w:t"), qn("w:delText"), qn("w:instrText"))
    breaks = (qn("w:tab"), qn("w:br"), qn("w:cr"))
    for part in parts:
        # adres hiperlacza (mailto:...) siedzi w .rels, poza tekstem akapitu
        for r in part.rels.values():
            if r.is_external:
                spans = find_spans(r.target_ref, enabled, words, ignore=ignore)
                if spans:
                    found.update(znalezione(r.target_ref, spans))
                    if not skan:
                        r._target = mask(r.target_ref, spans)
        for p in part.element.iter(qn("w:p")):
            # tabulator i zlamanie linii w runie = odstep miedzy slowami (inaczej
            # "Kowalski<tab>koniec" skleja sie w jedno slowo i lista slow go nie widzi);
            # w:tab w ustawieniach akapitu (pozycje tabulatorow) to nie tekst
            nodes = [n for n in p.iter(*tags, *breaks) if n.tag in tags or n.getparent().tag == qn("w:r")]
            texts = [(n.text or "") if n.tag in tags else " " for n in nodes]
            text = "".join(texts)
            spans = find_spans(text, enabled, words, ignore=ignore)
            if not spans:
                continue
            found.update(znalezione(text, spans))
            if skan:
                pos = 0
                for n, t in zip(nodes, texts):
                    a, b = pos, pos + len(t)
                    pos = b
                    if n.tag != qn("w:t") or not any(x < b and y > a for x, y, _ in spans):
                        continue
                    # granice zaznaczen wewnatrz tego fragmentu tekstu
                    ciecia = sorted({a, b} | {min(max(v, a), b) for x, y, _ in spans for v in (x, y)})
                    zaznaczone = [any(x <= c < y for x, y, _ in spans) for c in ciecia[:-1]]
                    zaznacz_docx(n, [(text[c:d], z) for c, d, z in zip(ciecia, ciecia[1:], zaznaczone)])
                continue
            masked, pos = mask(text, spans), 0
            # tekst akapitu bywa pociety na wiele "runow" - maska ma ta sama
            # dlugosc, wiec kazdy fragment dostaje swoj wycinek
            for n, t in zip(nodes, texts):
                if n.tag in tags:
                    n.text = masked[pos:pos + len(t)]
                pos += len(t)
    if skan:  # kopia robocza z zaznaczeniami - metadane bez zmian
        doc.save(dst)
        return found, 0
    cp = doc.core_properties
    for attr in ("author", "last_modified_by", "comments", "title",
                 "subject", "keywords", "category"):
        setattr(cp, attr, "")
    # app.xml (Firma, Menedzer) i wlasciwosci niestandardowe - Word dziala bez nich
    pkg_rels = doc.part.package.rels
    for rid, r in list(pkg_rels.items()):
        if r.reltype.endswith(("/extended-properties", "/custom-properties")):
            del pkg_rels[rid]
    doc.save(dst)
    return found, 0


# ------------------------------------------------------------------- PDF ----


def anon_pdf(src, dst, enabled, words, ignore=(), skan=False):
    import pymupdf

    doc = pymupdf.open(src)
    if doc.needs_pass:
        doc.close()
        raise ValueError("PDF jest zabezpieczony haslem - zdejmij haslo i sprobuj ponownie")
    stats, missed, has_text = Counter(), 0, False
    zolty, czerwony = (1, 0.9, 0), (1, 0, 0)

    def ramka_pdf(rect):  # skanowanie: obramowanie, tresc pod spodem zostaje
        a = page.add_rect_annot(rect)
        a.set_colors(stroke=czerwony)
        a.set_border(width=1.5)
        a.update()

    for page in doc:
        # pola formularza (wypelnione wnioski) nie sa czescia tekstu strony
        pola = []
        for w in page.widgets():
            v = w.field_value
            if isinstance(v, str) and v.strip():
                has_text = True
                pola.append(w.rect)
                spans = find_spans(v, enabled, words, ignore=ignore)
                if spans:
                    stats.update(znalezione(v, spans))
                    if skan:
                        ramka_pdf(w.rect)
                    else:
                        w.field_value = mask(v, spans)
                        w.update()
        text = page.get_text()
        found = False
        if not text.strip() and page.get_images():  # strona skanu - OCR
            has_text = True
            # piksele OCR -> wspolrzedne strony (takze strony obroconej)
            to_page = ~(page.rotation_matrix * pymupdf.Matrix(OCR_DPI / 72, OCR_DPI / 72))
            png = page.get_pixmap(dpi=OCR_DPI).tobytes("png")
            for cat, frag, boxes in ocr_hits(png, enabled, words, ignore):
                for b in boxes:
                    if skan:
                        ramka_pdf(pymupdf.Rect(b) * to_page)
                    else:
                        page.add_redact_annot(pymupdf.Rect(b) * to_page, fill=(0, 0, 0))
                stats[cat, frag] += 1
                found = True
        has_text = has_text or bool(text.strip())
        spans = find_spans(text, enabled, words, ignore=ignore)
        for (a, b, cat), (_, opis) in zip(spans, znalezione(text, spans)):
            frag = text[a:b]
            rects = []
            # fragment przelamany miedzy liniami szukamy linia po linii
            for piece in frag.split("\n"):
                if piece.strip():
                    rects += page.search_for(piece)
            if rects and any(all(r.intersects(p) for r in rects) for p in pola):
                continue  # wyglad pola formularza - juz policzone przy polach
            if rects:
                if skan:
                    zaz = page.add_highlight_annot(rects)
                    zaz.set_colors(stroke=zolty)
                    zaz.update()
                else:
                    for r in rects:
                        page.add_redact_annot(r, fill=(0, 0, 0))
                stats[cat, opis] += 1
                found = True
            else:
                missed += 1
        if found and not skan:
            page.apply_redactions()  # wycina tekst (i piksele obrazow) spod prostokatow
    if not has_text:
        doc.close()
        raise ValueError("PDF nie ma tekstu ani obrazow - nie ma czego anonimizowac")
    if not skan:  # kopia robocza ze skanowania - metadane bez zmian
        doc.set_metadata({})
        doc.del_xml_metadata()
    doc.save(dst, garbage=4, deflate=True)
    doc.close()
    return stats, missed


# ------------------------------------------------------------------ wsad ----

HANDLERS = {".docx": anon_docx, ".pdf": anon_pdf, **{e: anon_image for e in IMAGE_EXTS}}


def raport(found, wciecie="  "):
    """Linie raportu: kategoria z liczba, pod nia znalezione wartosci."""
    linie = []
    for cat, n in sorted(kategorie(found).items()):
        linie.append("%s%s: %d" % (wciecie, cat, n))
        for (c, t), k in sorted(found.items()):
            if c == cat:
                linie.append("%s    %s%s" % (wciecie, t, " (x%d)" % k if k > 1 else ""))
    return linie


def run_batch(inp, out_dir, enabled, words, log=print, ignore=(), skan=False):
    """skan=True: kopia z zaznaczeniami (_skan), nic nie jest usuwane.
    Zwraca Counter((kategoria, tekst)) ze wszystkich plikow."""
    if os.path.isfile(inp):
        files = [inp]
    else:
        files = sorted(os.path.join(inp, f) for f in os.listdir(inp)
                       if os.path.splitext(f)[1].lower() in HANDLERS)
    razem = Counter()
    if not files:
        log("Brak plikow DOCX/PDF/PNG/JPG w: %s" % inp)
        return razem
    os.makedirs(out_dir, exist_ok=True)
    for f in files:
        name, ext = os.path.splitext(os.path.basename(f))
        dst = os.path.join(out_dir, name + ("_skan" if skan else "_anonim") + ext.lower())
        log("%s: %s" % ("Skanuje" if skan else "Przetwarzam", os.path.basename(f)))
        try:
            found, missed = HANDLERS[ext.lower()](f, dst, enabled, words, ignore=ignore, skan=skan)
            razem.update(found)
            for linia in raport(found) or ["  nic nie znaleziono"]:
                log(linia)
            if missed:
                log("  UWAGA: %d znalezionych danych nie udalo sie zlokalizowac na stronie"
                    " - sprawdz plik recznie!" % missed)
            if any(c.endswith("(niepewny OCR)") for c, _ in found):
                log("  UWAGA: %s tez 11 cyfr z bledna suma (mozliwy blad OCR w PESEL)"
                    " - sprawdz wynik" % ("zaznaczono" if skan else "zakryto"))
            log("  OK -> %s" % dst)
        except Exception:
            log("BLAD: %s\n%s" % (os.path.basename(f), traceback.format_exc()))
    log("")
    log("=== Podsumowanie (%d plikow): %s ===" % (
        len(files), "znaleziono" if razem else "nic nie znaleziono"))
    for linia in raport(razem):
        log(linia)
    if skan:
        log("Skanowanie: dane zaznaczone w kopiach _skan, nic nie usunieto. Dopisz frazy do pominiecia"
            " albo slowa i kliknij Anonimizuj.")
    else:
        log("Zakonczono. Zawsze przejrzyj wynik przed publikacja.")
    return razem


# -------------------------------------------------------------------- GUI ----


def gui():
    import tkinter as tk
    from tkinter import filedialog, ttk, scrolledtext

    root = tk.Tk()
    root.title("Anonimizacja RODO (DOCX / PDF / PNG / JPG)")
    root.geometry("780x720")
    pad = dict(padx=6, pady=3)

    v_in = tk.StringVar(value=os.path.join(APP_DIR, "INPUT"))
    v_out = tk.StringVar(value=os.path.join(APP_DIR, "OUTPUT"))
    v_cats = {c: tk.BooleanVar(value=True) for c in CATEGORIES}

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)
    ttk.Label(f, text="Plik lub folder:").grid(row=0, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_in, width=60).grid(row=0, column=1, **pad)
    ttk.Button(f, text="Plik...", command=lambda: v_in.set(filedialog.askopenfilename(
        filetypes=[("DOCX / PDF / PNG / JPG", "*.docx *.pdf *.png *.jpg *.jpeg")]) or v_in.get())
               ).grid(row=0, column=2, **pad)
    ttk.Button(f, text="Folder...", command=lambda: v_in.set(
        filedialog.askdirectory() or v_in.get())).grid(row=0, column=3, **pad)
    ttk.Label(f, text="Folder wyjsciowy:").grid(row=1, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_out, width=60).grid(row=1, column=1, **pad)
    ttk.Button(f, text="Wybierz...", command=lambda: v_out.set(
        filedialog.askdirectory() or v_out.get())).grid(row=1, column=2, **pad)

    g = ttk.LabelFrame(root, text="Co szukac")
    g.pack(fill="x", **pad)
    for i, c in enumerate(CATEGORIES):
        ttk.Checkbutton(g, text=c, variable=v_cats[c]).grid(row=0, column=i, sticky="w", **pad)

    h = ttk.LabelFrame(root, text="Dodatkowe slowa do usuniecia (imiona, nazwiska, adresy) - jedno w linii")
    h.pack(fill="x", **pad)
    words_box = tk.Text(h, height=4)
    words_box.pack(fill="x", **pad)

    h2 = ttk.LabelFrame(root, text="Frazy do pominiecia - zostaja w dokumencie (np. NIP, telefon, e-mail urzedu) "
                                   "- jedna w linii")
    h2.pack(fill="x", **pad)
    ignore_box = tk.Text(h2, height=3)
    ignore_box.pack(fill="x", **pad)

    log_box = scrolledtext.ScrolledText(root, height=14)
    log_box.pack(fill="both", expand=True, **pad)

    def log(msg):
        def put():
            log_box.insert("end", str(msg) + "\n")
            log_box.see("end")
        root.after(0, put)

    b = ttk.Frame(root)
    b.pack(pady=6)
    btn_skan = ttk.Button(b, text="Skanuj (tylko zaznacz)")
    btn_skan.pack(side="left", padx=6)
    btn = ttk.Button(b, text="Anonimizuj")
    btn.pack(side="left", padx=6)

    def start(skan):
        inp, out = v_in.get().strip('" '), v_out.get().strip('" ')
        if not os.path.exists(inp):
            return log("Wskaz istniejacy plik lub folder.")
        if not out:
            return log("Wskaz folder wyjsciowy.")
        enabled = {c for c, v in v_cats.items() if v.get()}
        words = words_box.get("1.0", "end").splitlines()
        ignore = ignore_box.get("1.0", "end").splitlines()
        for x in (btn, btn_skan):
            x.config(state="disabled")
        log_box.delete("1.0", "end")

        def work():
            try:
                run_batch(inp, out, enabled, words, log, ignore=ignore, skan=skan)
            finally:
                root.after(0, lambda: [x.config(state="normal") for x in (btn, btn_skan)])

        threading.Thread(target=work, daemon=True).start()

    btn.config(command=lambda: start(False))
    btn_skan.config(command=lambda: start(True))
    if "--selftest" in sys.argv:
        root.after(200, root.destroy)
    import aktualizacja
    aktualizacja.start(root, "DawidBochno/Anonimizacja-RODO", "main", "anonimizacja.py")
    root.mainloop()


# --------------------------------------------------------------- selftest ----


def selftest():
    import tempfile
    import zipfile
    import docx
    import pymupdf

    assert pesel_ok("44051401359") and not pesel_ok("44051401358")
    assert nip_ok("123-456-32-18") and not nip_ok("1234563219")
    assert regon_ok("123456785") and regon_ok("12345678512347") and not regon_ok("123456789")
    assert dowod_ok("ABA300000") and not dowod_ok("ABA300001")
    assert iban_ok("PL61 1090 1014 0000 0712 1981 2874")
    assert iban_ok("61109010140000071219812874") and not iban_ok("61109010140000071219812875")

    t = ("PESEL 44051401359, NIP 123-456-32-18, kwota 12345678901, "
         "dowod ABA 300000, tel. +48 600 700 800, mail jan.kowalski@urzad.gov.pl, "
         "konto PL61 1090 1014 0000 0712 1981 2874, Jan Kowalski.")
    spans = find_spans(t, None, ["Jan Kowalski"])
    cats = Counter(s[2] for s in spans)
    assert cats == Counter({"PESEL": 1, "NIP": 1, "Dowód": 1, "Telefon": 1,
                            "E-mail": 1, "IBAN": 1, "Lista słów": 1}), cats
    m = mask(t, spans)
    assert len(m) == len(t) and "12345678901" in m and "44051401359" not in m
    assert "Kowalski." not in m and "kowalski@" not in m
    assert [s[2] for s in find_spans(t, {"E-mail"}, ())] == ["E-mail"]
    # dywizy i twarde spacje z Worda / PDF (U+00AD, U+2011, U+00A0)
    t2 = "NIP 123­456­32­18, tel. 600 100 200, 123‑456‑32‑18"
    assert [s[2] for s in find_spans(t2)] == ["NIP", "Telefon", "NIP"], find_spans(t2)

    # frazy do pominiecia: dane urzedu zostaja, takze w innym zapisie i w srodku frazy
    t3 = ("Urząd Gminy Kowalewo, NIP 123-456-32-18, sekretariat@urzad.gov.pl; wnioskodawca "
          "Jan Kowalski, jan@x.pl, ul. Kowalewo 5, NIP 1234563218")
    ign = ["1234563218", "SEKRETARIAT@urzad.gov.pl", "Urząd  Gminy Kowalewo", " "]
    spans = find_spans(t3, None, ["Kowalewo", "Jan Kowalski"], ignore=ign)
    assert znalezione(t3, spans) == [("Lista słów", "Jan Kowalski"), ("E-mail", "jan@x.pl"),
                                     ("Lista słów", "Kowalewo")], znalezione(t3, spans)
    assert [s[2] for s in find_spans("tel. 600 700 800", ignore=["600"])] == []  # fraza w srodku numeru
    assert znalezione("PESEL 4405\n1401359", [(6, 18, "PESEL")]) == [("PESEL", "4405 1401359")]
    assert raport(Counter({("PESEL", "44051401359"): 2, ("E-mail", "a@b.pl"): 1})) == [
        "  E-mail: 1", "      a@b.pl", "  PESEL: 2", "      44051401359 (x2)"]

    tmp = tempfile.mkdtemp()
    # DOCX: PESEL pociety na dwa runy + stopka
    d = docx.Document()
    p = d.add_paragraph("PESEL: 4405")
    p.add_run("1401359 koniec")
    d.sections[0].footer.paragraphs[0].text = "kontakt: a@b.pl"
    d.core_properties.author = "Jan Kowalski"
    # hiperlacze mailto: adres jest w .rels, nie w tekscie
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import parse_xml
    rid = d.part.relate_to("mailto:jan.kowalski@urzad.gov.pl", RT.HYPERLINK,
                           is_external=True)
    d.add_paragraph()._p.append(parse_xml(
        '<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
        ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        ' r:id="%s"><w:r><w:t>napisz</w:t></w:r></w:hyperlink>' % rid))
    # pole Firma w docProps/app.xml
    app = next(p for p in d.part.package.iter_parts() if p.partname == "/docProps/app.xml")
    app._blob = app.blob.replace(b"</Properties>", b"<Company>Kowalski</Company></Properties>")
    src = os.path.join(tmp, "t.docx")
    d.save(src)
    assert b"Kowalski</Company>" in zipfile.ZipFile(src).read("docProps/app.xml")
    anon_docx(src, os.path.join(tmp, "o.docx"), None, ())
    with zipfile.ZipFile(os.path.join(tmp, "o.docx")) as z:
        raw = b"".join(z.read(n) for n in z.namelist())
    assert b"kowalski" not in raw.lower(), "wyciek w hiperlaczu lub wlasciwosciach"
    o = docx.Document(os.path.join(tmp, "o.docx"))
    assert o.paragraphs[0].text == "PESEL: *********** koniec", o.paragraphs[0].text
    assert "a@b.pl" not in o.sections[0].footer.paragraphs[0].text
    assert o.core_properties.author == ""

    # DOCX skanowanie: tekst bez zmian, dane na zolto (takze pociete na runy), metadane zostaja
    d = docx.Document()
    p = d.add_paragraph("PESEL: 4405")
    p.add_run("1401359 i Jan ").bold = True
    r = p.add_run("Kowalski")
    r.add_tab()
    r.add_text("koniec")
    d.add_paragraph("NIP urzedu 123-456-32-18")
    d.core_properties.author = "Jan Kowalski"
    d.save(src)
    found, _ = anon_docx(src, os.path.join(tmp, "s.docx"), None, ["Jan Kowalski"],
                         ignore=["123-456-32-18"], skan=True)
    assert found == Counter({("PESEL", "44051401359"): 1, ("Lista słów", "Jan Kowalski"): 1}), found
    o = docx.Document(os.path.join(tmp, "s.docx"))
    p = o.paragraphs[0]
    assert p.text == "PESEL: 44051401359 i Jan Kowalski\tkoniec", repr(p.text)
    assert [r.text for r in p.runs if r.font.highlight_color] == ["4405", "1401359", "Jan ", "Kowalski"]
    assert [r.text for r in p.runs if r.bold] == ["1401359", " i ", "Jan "]
    assert not any(r.font.highlight_color for r in o.paragraphs[1].runs)
    assert o.core_properties.author == "Jan Kowalski"
    # anonimizacja z fraza do pominiecia
    anon_docx(src, os.path.join(tmp, "o2.docx"), None, ["Kowalski"], ignore=["1234563218"])
    o = docx.Document(os.path.join(tmp, "o2.docx"))
    assert o.paragraphs[1].text == "NIP urzedu 123-456-32-18"
    assert o.paragraphs[0].text == "PESEL: *********** i Jan ********	koniec", repr(o.paragraphs[0].text)

    # PDF: tekst ma faktycznie zniknac z pliku
    pdoc = pymupdf.open()
    pdoc.new_page().insert_text((72, 72), "Wnioskodawca PESEL 44051401359 zostaje")
    pdoc.set_metadata({"author": "Jan Kowalski"})
    pdoc.save(os.path.join(tmp, "t.pdf"))
    # skanowanie: podswietlenie, tekst i metadane zostaja
    stats, missed = anon_pdf(os.path.join(tmp, "t.pdf"), os.path.join(tmp, "s.pdf"), None, (), skan=True)
    assert stats == Counter({("PESEL", "44051401359"): 1}) and missed == 0, stats
    with pymupdf.open(os.path.join(tmp, "s.pdf")) as out:
        assert "44051401359" in out[0].get_text() and out.metadata.get("author") == "Jan Kowalski"
        assert [a.type[1] for a in out[0].annots()] == ["Highlight"]
    stats, _ = anon_pdf(os.path.join(tmp, "t.pdf"), os.path.join(tmp, "i.pdf"), None, (), ignore=["44051401359"])
    assert not stats
    stats, missed = anon_pdf(os.path.join(tmp, "t.pdf"), os.path.join(tmp, "o.pdf"), None, ())
    assert kategorie(stats)["PESEL"] == 1 and missed == 0
    out = pymupdf.open(os.path.join(tmp, "o.pdf"))
    txt = out[0].get_text()
    assert "44051401359" not in txt and "Wnioskodawca" in txt, txt
    assert not out.metadata.get("author")
    out.close()

    # PDF: wypelnione pole formularza
    pdoc = pymupdf.open()
    page = pdoc.new_page()
    w = pymupdf.Widget()
    w.field_type, w.field_name = pymupdf.PDF_WIDGET_TYPE_TEXT, "pesel"
    w.rect, w.field_value = pymupdf.Rect(72, 100, 300, 120), "44051401359"
    page.add_widget(w)
    pdoc.save(os.path.join(tmp, "f.pdf"))
    stats, _ = anon_pdf(os.path.join(tmp, "f.pdf"), os.path.join(tmp, "fo.pdf"), None, ())
    assert kategorie(stats)["PESEL"] == 1, stats
    out = pymupdf.open(os.path.join(tmp, "fo.pdf"))
    vals = [x.field_value for x in out[0].widgets()]
    assert vals == ["***********"], vals
    assert "44051401359" not in out[0].get_text()
    out.close()
    stats, _ = anon_pdf(os.path.join(tmp, "f.pdf"), os.path.join(tmp, "fs.pdf"), None, (), skan=True)
    assert stats == Counter({("PESEL", "44051401359"): 1}), stats  # pole liczone raz, nie tez z wygladu
    with pymupdf.open(os.path.join(tmp, "fs.pdf")) as out:
        assert [x.field_value for x in out[0].widgets()] == ["44051401359"]
        assert [a.type[1] for a in out[0].annots()] == ["Square"]

    # wsad: skanowanie folderu - nazwy _skan, raport z wartosciami, podsumowanie
    wsad = os.path.join(tmp, "wsad")
    os.makedirs(wsad)
    for n in ("t.pdf", "t.docx"):
        shutil.copy(os.path.join(tmp, n), wsad)
    lines = []
    razem = run_batch(wsad, os.path.join(tmp, "wynik"), None, ["Jan Kowalski"], lines.append, skan=True)
    assert sorted(os.listdir(os.path.join(tmp, "wynik"))) == ["t_skan.docx", "t_skan.pdf"]
    assert razem[("PESEL", "44051401359")] == 2, razem
    assert "    44051401359 (x2)" in "\n".join(lines) and any("Podsumowanie (2 plikow)" in x for x in lines)

    # OCR: PESEL z bledna suma tylko w tekscie z OCR
    assert find_spans("nr 44051401358") == []
    assert [s[2] for s in find_spans("nr 44051401358", {"PESEL"}, (), ocr=True)] == [
        "PESEL (niepewny OCR)"]
    assert find_spans("nr 44051401358", {"NIP"}, (), ocr=True) == []
    if find_tesseract():
        selftest_ocr(tmp)
    elif os.environ.get("CI"):
        raise AssertionError("CI: brak Tesseracta")
    else:
        print("UWAGA: brak Tesseracta - test skanow i zdjec pominiety")
    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK")


def selftest_ocr(tmp):
    import pymupdf

    def scan_pix(rot=0):
        d = pymupdf.open()
        p = d.new_page(width=500, height=200)
        p.insert_text((20, 60), "Wnioskodawca Jan Kowalski", fontsize=20)
        p.insert_text((20, 110), "PESEL 44051401359", fontsize=20)
        p.insert_text((20, 160), "Sygnatura 4321", fontsize=20)
        return p.get_pixmap(matrix=pymupdf.Matrix(3, 3).prerotate(rot))

    def ocr_text(png):
        found = ocr_hits(png, None, ["Wnioskodawca", "Sygnatura", "Kowalski"])
        return Counter(c for c, _, _ in found)

    # zdjecie z telefonu: zapisane bokiem + EXIF (obrot 6 i pole Artist)
    jpg = scan_pix(rot=-90).tobytes("jpg")
    exif = (b"Exif\x00\x00MM\x00\x2a\x00\x00\x00\x08\x00\x02"
            b"\x01\x12\x00\x03\x00\x00\x00\x01\x00\x06\x00\x00"
            b"\x01\x3b\x00\x02\x00\x00\x00\x04KOW\x00\x00\x00\x00\x00")
    src = os.path.join(tmp, "foto.jpg")
    with open(src, "wb") as f:
        f.write(jpg[:2] + b"\xff\xe1" + (len(exif) + 2).to_bytes(2, "big") + exif + jpg[2:])
    stats, _ = anon_image(src, os.path.join(tmp, "foto_o.jpg"), None, ["Kowalski"])
    assert stats == Counter({("PESEL", "44051401359"): 1, ("Lista słów", "Kowalski"): 1}), stats
    # skanowanie: ramki, dane dalej czytelne; fraza pominieta nie jest zaznaczana
    stats, _ = anon_image(src, os.path.join(tmp, "foto_s.jpg"), None, ["Kowalski"],
                          ignore=["Kowalski"], skan=True)
    assert stats == Counter({("PESEL", "44051401359"): 1}), stats
    with pymupdf.open(os.path.join(tmp, "foto_s.jpg")) as d:
        left = ocr_text(d[0].get_pixmap(dpi=OCR_DPI).tobytes("png"))
    assert left["PESEL"] == 1, left
    raw = open(os.path.join(tmp, "foto_o.jpg"), "rb").read()
    assert b"Exif" not in raw and b"KOW" not in raw, "metadane zdjecia zostaly"
    # po anonimizacji OCR nie moze juz znalezc danych, a reszta tekstu zostaje
    with pymupdf.open(os.path.join(tmp, "foto_o.jpg")) as d:
        left = ocr_text(d[0].get_pixmap(dpi=OCR_DPI).tobytes("png"))
    assert left == Counter({"Lista słów": 2}), left  # Wnioskodawca + Sygnatura, prosto

    # skan PDF (strona obrocona o 90 stopni, jak ze skanera)
    d = pymupdf.open()
    page = d.new_page(width=200, height=500)  # obraz lezy bokiem, /Rotate go prostuje
    page.insert_image(page.rect, pixmap=scan_pix(rot=-90))
    page.set_rotation(90)
    d.save(os.path.join(tmp, "skan.pdf"))
    stats, missed = anon_pdf(os.path.join(tmp, "skan.pdf"), os.path.join(tmp, "skan_o.pdf"),
                             None, ["Kowalski"])
    assert kategorie(stats) == Counter({"PESEL": 1, "Lista słów": 1}) and not missed, stats
    with pymupdf.open(os.path.join(tmp, "skan_o.pdf")) as d:
        left = ocr_text(d[0].get_pixmap(dpi=OCR_DPI).tobytes("png"))
    assert left == Counter({"Lista słów": 2}), f"skan: dane widoczne po anonimizacji: {left}"


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        if "--gui" in sys.argv:
            gui()
    else:
        gui()
